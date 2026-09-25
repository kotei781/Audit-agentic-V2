"""
report_exporter.py
-------------------
TẦNG XUẤT BÁO CÁO — Word (.docx) và JSON.

CHỐT CHẶN NGHIỆP VỤ: mọi hàm xuất đều gọi `_ensure_reviewed()` trước. Báo cáo
chỉ được phép rời khỏi hệ thống sau khi TOÀN BỘ finding đã có quyết định của
con người. Đây là ranh giới cứng của mô hình Human-in-the-Loop — đặt ở tầng
xuất file chứ không phải ở tầng UI, để không UI nào lách qua được.

Hàm `build_docx_bytes()` trả về bytes trong bộ nhớ (BytesIO) thay vì ghi ra
đĩa — cần thiết cho Streamlit Cloud, nơi filesystem là ephemeral và file ghi
xuống sẽ biến mất khi container restart.
"""

from __future__ import annotations

import io
from datetime import timezone
from pathlib import Path
from typing import Optional, Union

import config
from schemas import AuditReport, HumanDecision, RiskLevel, ViolationStatus

logger = config.get_logger("report_exporter")

try:
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor
except ImportError:  # pragma: no cover
    Document = None


class ExportError(RuntimeError):
    """Lỗi khi xuất báo cáo."""


class ReviewIncompleteError(ExportError):
    """Báo cáo chưa được kiểm duyệt đầy đủ — không được phép xuất."""


# Bảng màu theo mức rủi ro (hex, không có dấu #).
_RISK_COLORS = {
    RiskLevel.HIGH: "C0392B",
    RiskLevel.MEDIUM: "D68910",
    RiskLevel.LOW: "1E8449",
}

_RISK_LABELS = {
    RiskLevel.HIGH: "RỦI RO CAO",
    RiskLevel.MEDIUM: "RỦI RO TRUNG BÌNH",
    RiskLevel.LOW: "RỦI RO THẤP",
}

_STATUS_LABELS = {
    ViolationStatus.YES: "CÓ VI PHẠM",
    ViolationStatus.NO: "KHÔNG VI PHẠM",
    ViolationStatus.UNCERTAIN: "CHƯA ĐỦ CĂN CỨ",
}


# ============================================================
# TIỆN ÍCH
# ============================================================
def _ensure_reviewed(report: AuditReport) -> None:
    """Chặn xuất báo cáo khi còn finding chưa có quyết định của con người."""
    pending = report.pending_findings()
    if pending:
        raise ReviewIncompleteError(
            f"Không thể xuất báo cáo: còn {len(pending)} phát hiện chưa được con "
            "người xác nhận. Hoàn tất bước Human-in-the-Loop trước đã."
        )


def _local_time(report: AuditReport) -> str:
    """Hiển thị thời gian tạo báo cáo theo giờ địa phương kèm mốc UTC."""
    generated = report.generated_at
    if generated.tzinfo is None:
        generated = generated.replace(tzinfo=timezone.utc)
    local = generated.astimezone()
    return f"{local.strftime('%d/%m/%Y %H:%M:%S %Z')} (UTC: {generated.strftime('%Y-%m-%d %H:%M:%S')})"


def _shade_cell(cell, hex_color: str) -> None:
    """Tô nền một ô bảng (python-docx không có API sẵn, phải chạm vào XML)."""
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), hex_color)
    cell._tc.get_or_add_tcPr().append(shading)


def _add_kv_table(document, rows, col_widths=(2.0, 4.3)) -> None:
    """Thêm bảng 2 cột dạng nhãn - giá trị."""
    table = document.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    for label, value in rows:
        cells = table.add_row().cells
        cells[0].width = Inches(col_widths[0])
        cells[1].width = Inches(col_widths[1])

        run = cells[0].paragraphs[0].add_run(str(label))
        run.bold = True
        cells[1].paragraphs[0].add_run(str(value))
        _shade_cell(cells[0], "F2F3F4")


def _default_filename(report: AuditReport, extension: str) -> str:
    stamp = report.generated_at.strftime("%Y%m%d_%H%M%S")
    return f"audit_report_{stamp}_{report.report_id[:8]}.{extension}"


# ============================================================
# XUẤT WORD
# ============================================================
def build_docx_bytes(report: AuditReport, reviewer_id: Optional[str] = None) -> bytes:
    """
    Dựng báo cáo kiểm toán dạng .docx và trả về bytes (không ghi ra đĩa).

    Cấu trúc báo cáo:
        1. Trang tiêu đề + thông tin phiên kiểm toán (audit trail)
        2. Bảng tóm tắt kết quả
        3. Chi tiết từng phát hiện (kèm suy luận, trích dẫn, quyết định người duyệt)
        4. Ghi chú giới hạn & khuyến cáo sử dụng
    """
    if Document is None:
        raise ExportError("Chưa cài python-docx. Chạy: pip install python-docx")

    _ensure_reviewed(report)

    document = Document()

    # ---- Metadata của file Word ----
    properties = document.core_properties
    properties.title = f"Báo cáo kiểm toán {report.report_id}"
    properties.author = reviewer_id or report.findings[0].reviewer_id if report.findings else "Agentic AI Audit System"
    properties.comments = f"Sinh bởi Agentic AI Audit System — prompt {report.prompt_version}"

    # ---- 1. Tiêu đề ----
    heading = document.add_heading("BÁO CÁO KIỂM TOÁN CHỨNG TỪ TÀI CHÍNH", level=0)
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("Agentic AI Audit System — Kết quả đã qua kiểm duyệt của con người")
    run.italic = True
    run.font.size = Pt(10)

    if report.hitl_bypassed:
        warning = document.add_paragraph()
        warning.alignment = WD_ALIGN_PARAGRAPH.CENTER
        warning_run = warning.add_run(
            "⚠ BÁO CÁO NÀY ĐƯỢC TẠO Ở CHẾ ĐỘ TỰ ĐỘNG PHÊ DUYỆT — "
            "KHÔNG CÓ GIÁ TRỊ SỬ DỤNG CHÍNH THỨC"
        )
        warning_run.bold = True
        warning_run.font.color.rgb = RGBColor(0xC0, 0x39, 0x2B)

    # ---- 2. Thông tin phiên kiểm toán ----
    document.add_heading("1. Thông tin phiên kiểm toán", level=1)

    law_files = sorted({ref.filename for ref in report.law_sources})

    # [UPDATED] Xử lý hiển thị trạng thái truncation trong bảng thông tin
    truncation_status = "Không" if not report.document_truncated else (
        f"Có ({report.truncation_details or 'Chi tiết không xác định'})"
    )

    _add_kv_table(
        document,
        [
            ("Mã báo cáo", report.report_id),
            ("Thời điểm tạo", _local_time(report)),
            ("Chứng từ kiểm toán", report.document_source),
            ("SHA-256 chứng từ", report.document_sha256 or "(không ghi nhận)"),
            ("Chế độ truy hồi luật", report.retrieval_mode),
            ("Văn bản pháp lý tham chiếu", ", ".join(law_files) if law_files else "Toàn bộ kho luật"),
            ("Mô hình AI", f"{report.provider_used} / {report.model_used}"),
            ("Mô hình embedding", report.embedding_model or "(không dùng)"),
            ("Phiên bản prompt", report.prompt_version),
            ("Temperature", f"{report.temperature}"),
            ("Tài liệu bị cắt bớt", truncation_status), # [UPDATED]
            ("Người kiểm duyệt", reviewer_id or _collect_reviewers(report)),
        ],
    )

    # ---- 3. Tóm tắt kết quả ----
    document.add_heading("2. Tóm tắt kết quả", level=1)
    summary = report.summary()

    table = document.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    headers = ["Chỉ tiêu", "Số lượng", "Chỉ tiêu", "Số lượng"]
    for index, text in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.paragraphs[0].add_run(text).bold = True
        _shade_cell(cell, "D6EAF8")

    summary_pairs = [
        ("Tổng số phát hiện", summary["total"]),
        ("Kết luận CÓ vi phạm", summary["violation"]),
        ("Chưa đủ căn cứ (UNCERTAIN)", summary["uncertain"]),
        ("Không vi phạm", summary["compliant"]),
        ("Rủi ro cao", summary["high_risk"]),
        ("Trích dẫn không kiểm chứng được", summary["grounding_failed"]),
        ("Đã phê duyệt", summary["approved"]),
        ("Đã từ chối", summary["rejected"]),
    ]

    for index in range(0, len(summary_pairs), 2):
        cells = table.add_row().cells
        cells[0].paragraphs[0].add_run(summary_pairs[index][0])
        cells[1].paragraphs[0].add_run(str(summary_pairs[index][1])).bold = True
        if index + 1 < len(summary_pairs):
            cells[2].paragraphs[0].add_run(summary_pairs[index + 1][0])
            cells[3].paragraphs[0].add_run(str(summary_pairs[index + 1][1])).bold = True

    # ---- 4. Chi tiết phát hiện ----
    document.add_heading("3. Chi tiết các phát hiện", level=1)

    if not report.findings:
        document.add_paragraph(
            "Không có phát hiện nào. AI đánh giá chứng từ phù hợp với các quy "
            "định trong phạm vi văn bản pháp lý đã cung cấp."
        )
    else:
        approved_first = report.approved_findings() + report.rejected_findings()
        for index, finding in enumerate(approved_first, start=1):
            _write_finding(document, index, finding)

    # ---- 5. Khuyến cáo ----
    document.add_heading("4. Phạm vi và giới hạn", level=1)
    for line in (
        "Báo cáo này được hỗ trợ bởi mô hình ngôn ngữ lớn và CHỈ có giá trị "
        "tham khảo phục vụ công tác rà soát. Mọi kết luận đều phải được kiểm "
        "toán viên có thẩm quyền xác nhận trước khi sử dụng cho mục đích "
        "kê khai, quyết toán hoặc pháp lý.",
        "Phạm vi đối chiếu bị giới hạn trong các văn bản pháp luật đã được nạp "
        "vào kho của hệ thống tại thời điểm chạy. Điều khoản nằm ngoài kho — "
        "kể cả khi đang có hiệu lực — sẽ không được xét đến.",
        "Chỉ số confidence_score là tự đánh giá của mô hình, KHÔNG phải xác "
        "suất đã hiệu chuẩn thống kê, và không nên dùng làm ngưỡng tự động hóa.",
        "Các phát hiện có ghi chú 'trích dẫn không kiểm chứng được' là trường "
        "hợp mô hình đưa ra dẫn chứng không tìm thấy trong chứng từ gốc; cần "
        "đặc biệt thận trọng với nhóm này.",
    ):
        paragraph = document.add_paragraph(line, style="List Bullet")
        paragraph.paragraph_format.space_after = Pt(4)

    buffer = io.BytesIO()
    document.save(buffer)
    buffer.seek(0)
    logger.info("Đã dựng báo cáo Word cho report_id=%s", report.report_id)
    return buffer.getvalue()


def _collect_reviewers(report: AuditReport) -> str:
    reviewers = sorted({f.reviewer_id for f in report.findings if f.reviewer_id})
    return ", ".join(reviewers) if reviewers else "(không ghi nhận)"


def _write_finding(document, index: int, finding) -> None:
    """Ghi chi tiết một phát hiện vào tài liệu Word."""
    risk = finding.risk_level
    document.add_heading(
        f"3.{index}. Phát hiện #{index} — {_STATUS_LABELS.get(finding.has_violation, '')}",
        level=2,
    )

    # Thanh nhãn mức rủi ro
    label_table = document.add_table(rows=1, cols=1)
    label_table.style = "Table Grid"
    cell = label_table.rows[0].cells[0]
    run = cell.paragraphs[0].add_run(
        f"{_RISK_LABELS[risk]}  |  Độ tin cậy của AI: {finding.confidence_score:.2f}"
        f"  |  Quyết định: {_decision_label(finding.human_decision)}"
    )
    run.bold = True
    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    _shade_cell(cell, _RISK_COLORS[risk])

    rows = [
        ("Trích dẫn từ chứng từ", finding.input_citation),
        ("Căn cứ pháp lý", finding.rule_citation),
        ("Suy luận của AI", finding.reasoning_steps),
        ("Giải thích", finding.explanation),
        ("Kiểm chứng trích dẫn", _grounding_label(finding)),
        ("Hành động đề xuất", finding.recommended_action.value),
        ("Ghi chú của người duyệt", finding.human_note or "(không có)"),
        ("Người duyệt / thời điểm", _reviewer_label(finding)),
    ]
    if finding.downgraded_reason:
        rows.insert(5, ("Hệ thống tự hạ cấp", finding.downgraded_reason))

    _add_kv_table(document, rows)
    document.add_paragraph()


def _decision_label(decision: Optional[HumanDecision]) -> str:
    if decision == HumanDecision.APPROVED:
        return "ĐÃ PHÊ DUYỆT"
    if decision == HumanDecision.REJECTED:
        return "ĐÃ TỪ CHỐI"
    return "CHƯA XỬ LÝ"


def _grounding_label(finding) -> str:
    if finding.grounding_verified is None:
        return "Chưa kiểm chứng"
    if finding.grounding_verified:
        return f"Đạt — trích dẫn khớp tài liệu gốc (điểm {finding.grounding_score:.2f})"
    return (
        f"KHÔNG ĐẠT — không tìm thấy trích dẫn trong chứng từ gốc "
        f"(điểm {finding.grounding_score:.2f})"
    )


def _reviewer_label(finding) -> str:
    if not finding.reviewed_at:
        return "(chưa ghi nhận)"
    stamp = finding.reviewed_at
    if stamp.tzinfo is None:
        stamp = stamp.replace(tzinfo=timezone.utc)
    return f"{finding.reviewer_id or 'ẩn danh'} — {stamp.astimezone().strftime('%d/%m/%Y %H:%M')}"


# ============================================================
# GHI RA ĐĨA
# ============================================================
def export_report_to_docx(
    report: AuditReport,
    output_path: Union[str, Path, None] = None,
    reviewer_id: Optional[str] = None,
) -> Path:
    """Xuất báo cáo .docx ra đĩa và trả về đường dẫn tuyệt đối."""
    data = build_docx_bytes(report, reviewer_id=reviewer_id)
    path = Path(output_path) if output_path else config.OUTPUT_DIR / _default_filename(report, "docx")
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(data)
    logger.info("Đã xuất báo cáo Word: %s", path)
    return path.resolve()


def build_json_bytes(report: AuditReport) -> bytes:
    """Trả về báo cáo dạng JSON (bytes) — dùng cho nút tải trên UI."""
    _ensure_reviewed(report)
    return report.model_dump_json(indent=2).encode("utf-8")


def export_report_to_json(
    report: AuditReport,
    output_path: Union[str, Path, None] = None,
) -> Path:
    """Xuất báo cáo JSON ra đĩa (định dạng máy đọc được, dùng để lưu trữ/đối soát)."""
    data = build_json_bytes(report)
    path = Path(output_path) if output_path else config.OUTPUT_DIR / _default_filename(report, "json")
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(data)
    logger.info("Đã xuất báo cáo JSON: %s", path)
    return path.resolve()


def export_final_report(
    report: AuditReport,
    reviewer_id: Optional[str] = None,
) -> dict:
    """
    Xuất đồng thời cả hai định dạng — dùng cho luồng CLI.

    Returns:
        {"docx": Path, "json": Path}
    """
    return {
        "docx": export_report_to_docx(report, reviewer_id=reviewer_id),
        "json": export_report_to_json(report),
    }


def suggest_filename(report: AuditReport, extension: str = "docx") -> str:
    """Tên file gợi ý cho nút download trên Streamlit (an toàn trên mọi OS)."""
    return _default_filename(report, extension)
